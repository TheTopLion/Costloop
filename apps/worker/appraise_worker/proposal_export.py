"""Deterministic, client-safe XLSX proposal generation and verification."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import xlsxwriter
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field


class ProposalLine(BaseModel):
    model_config = ConfigDict(extra="ignore")

    position: str = ""
    description: str = Field(min_length=1, max_length=4000)
    quantity: str = ""
    unit: str = ""
    unit_price: str = ""
    line_total: str = ""


class ProposalPayload(BaseModel):
    model_config = ConfigDict(extra="forbid")

    seller_name: str = Field(min_length=1, max_length=200)
    seller_legal_name: str = Field(default="", max_length=300)
    client_name: str = Field(default="", max_length=300)
    project_name: str = Field(min_length=1, max_length=300)
    proposal_number: str = Field(min_length=1, max_length=160)
    version: int = Field(ge=1)
    currency: str = Field(pattern=r"^[A-Z]{3}$")
    created_at: str = ""
    validity_days: int | None = Field(default=None, ge=0, le=3650)
    notes: str = Field(default="", max_length=8000)
    commercial_terms: str = Field(default="", max_length=8000)
    subtotal: str
    markup_amount: str
    tax_amount: str
    grand_total: str
    lines: list[ProposalLine] = Field(max_length=5000)


def _number(value: str) -> float | None:
    try:
        return float(value) if value.strip() else None
    except ValueError:
        return None


def build_workbook(raw: dict[str, Any]) -> bytes:
    proposal = ProposalPayload.model_validate(raw)
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    workbook.set_properties({
        "title": f"Proposal {proposal.proposal_number}",
        "subject": proposal.project_name,
        "author": proposal.seller_name,
        "company": proposal.seller_name,
        "comments": "Generated from an immutable Appraise proposal version.",
    })
    sheet = workbook.add_worksheet("Proposal")
    navy, green, pale, gray = "#17352C", "#257A5A", "#EAF3EE", "#66736D"
    title = workbook.add_format({"bold": True, "font_size": 22, "font_color": navy, "align": "left", "valign": "vcenter"})
    meta_label = workbook.add_format({"bold": True, "font_color": gray, "font_size": 9})
    meta_value = workbook.add_format({"font_color": navy, "font_size": 10})
    header = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": navy, "border": 0, "align": "center", "valign": "vcenter"})
    text = workbook.add_format({"font_color": navy, "valign": "top", "text_wrap": True, "border": 1, "border_color": "#D8E0DC"})
    center = workbook.add_format({"font_color": navy, "align": "center", "valign": "top", "border": 1, "border_color": "#D8E0DC"})
    number = workbook.add_format({"font_color": navy, "num_format": "#,##0.00", "align": "right", "border": 1, "border_color": "#D8E0DC"})
    total_label = workbook.add_format({"bold": True, "font_color": navy, "align": "right", "top": 1, "top_color": green})
    total_value = workbook.add_format({"bold": True, "font_color": navy, "num_format": "#,##0.00", "align": "right", "top": 1, "top_color": green})
    grand_label = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": green, "align": "right"})
    grand_value = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": green, "num_format": "#,##0.00", "align": "right"})
    note_format = workbook.add_format({"font_color": gray, "bg_color": pale, "text_wrap": True, "valign": "top"})

    sheet.hide_gridlines(2)
    sheet.set_landscape()
    sheet.set_paper(9)
    sheet.fit_to_pages(1, 0)
    sheet.set_margins(0.35, 0.35, 0.55, 0.55)
    sheet.set_header("&RAppraise proposal")
    sheet.set_footer("&LConfidential&CPage &P of &N&R" + proposal.proposal_number)
    sheet.set_column("A:A", 12)
    sheet.set_column("B:B", 54)
    sheet.set_column("C:C", 13)
    sheet.set_column("D:D", 11)
    sheet.set_column("E:F", 18)
    sheet.set_row(0, 34)
    sheet.merge_range("A1:F1", "COMMERCIAL PROPOSAL", title)
    sheet.write("A3", "Seller", meta_label)
    sheet.merge_range("B3:C3", proposal.seller_legal_name or proposal.seller_name, meta_value)
    sheet.write("D3", "Proposal", meta_label)
    sheet.merge_range("E3:F3", f"{proposal.proposal_number} · v{proposal.version}", meta_value)
    sheet.write("A4", "Client", meta_label)
    sheet.merge_range("B4:C4", proposal.client_name or "—", meta_value)
    sheet.write("D4", "Currency", meta_label)
    sheet.merge_range("E4:F4", proposal.currency, meta_value)
    sheet.write("A5", "Project", meta_label)
    sheet.merge_range("B5:C5", proposal.project_name, meta_value)
    sheet.write("D5", "Date / validity", meta_label)
    validity = f" · valid {proposal.validity_days} days" if proposal.validity_days is not None else ""
    sheet.merge_range("E5:F5", f"{proposal.created_at[:10]}{validity}".strip(), meta_value)

    header_row = 7
    headers = ["Position", "Description", "Quantity", "Unit", "Unit price", "Line total"]
    for column, heading in enumerate(headers):
        sheet.write(header_row, column, heading, header)
    sheet.set_row(header_row, 26)
    first_data_row = header_row + 1
    for offset, line in enumerate(proposal.lines):
        row = first_data_row + offset
        sheet.write(row, 0, line.position, center)
        sheet.write(row, 1, line.description, text)
        quantity = _number(line.quantity)
        sheet.write_number(row, 2, quantity, number) if quantity is not None else sheet.write(row, 2, line.quantity, center)
        sheet.write(row, 3, line.unit, center)
        unit_price = _number(line.unit_price)
        sheet.write_number(row, 4, unit_price, number) if unit_price is not None else sheet.write_blank(row, 4, None, number)
        line_total = _number(line.line_total)
        sheet.write_number(row, 5, line_total, number) if line_total is not None else sheet.write_blank(row, 5, None, number)
    last_data_row = max(first_data_row, first_data_row + len(proposal.lines) - 1)
    sheet.autofilter(header_row, 0, last_data_row, 5)
    sheet.freeze_panes(first_data_row, 0)

    totals_row = last_data_row + 2
    totals = [("Subtotal", proposal.subtotal, total_label, total_value)]
    if proposal.markup_amount.strip():
        totals.append(("Pricing adjustment", proposal.markup_amount, total_label, total_value))
    totals.extend([
        ("Tax", proposal.tax_amount, total_label, total_value),
        ("Grand total", proposal.grand_total, grand_label, grand_value),
    ])
    for offset, (label, persisted, label_format, value_format) in enumerate(totals):
        row = totals_row + offset
        sheet.merge_range(row, 3, row, 4, label, label_format)
        numeric = _number(persisted)
        sheet.write_number(row, 5, numeric, value_format) if numeric is not None else sheet.write(row, 5, persisted, value_format)
    sheet.write(totals_row + len(totals) + 1, 0, "Notes and commercial terms", meta_label)
    sheet.merge_range(totals_row + len(totals) + 2, 0, totals_row + len(totals) + 4, 5,
                      "\n".join(value for value in (proposal.notes, proposal.commercial_terms) if value) or "—", note_format)
    sheet.print_area(0, 0, totals_row + len(totals) + 4, 5)
    sheet.repeat_rows(header_row)
    workbook.close()
    result = output.getvalue()
    verify_workbook(result, proposal)
    return result


def verify_workbook(content: bytes, proposal: ProposalPayload) -> None:
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    if workbook.sheetnames != ["Proposal"]:
        raise ValueError("Unexpected proposal workbook sheets")
    sheet = workbook["Proposal"]
    if sheet["A1"].value != "COMMERCIAL PROPOSAL" or sheet["A8"].value != "Position":
        raise ValueError("Proposal workbook structure verification failed")
    if sheet.freeze_panes != "A9" or not sheet.print_area:
        raise ValueError("Proposal workbook page layout verification failed")
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    if proposal.grand_total not in {str(value) for value in values}:
        numeric_total = _number(proposal.grand_total)
        if numeric_total is None or not any(isinstance(value, (int, float)) and abs(value - numeric_total) < 1e-7 for value in values):
            raise ValueError("Persisted grand total is missing from proposal workbook")
    forbidden = ("supplier_id", "quote_line_id", "confidence", "hard_conflicts")
    flattened = " ".join(str(value) for value in values if value is not None).lower()
    if any(term in flattened for term in forbidden):
        raise ValueError("Internal-only fields leaked into proposal workbook")
