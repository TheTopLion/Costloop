"""Bounded, deterministic source readers. No extracted value becomes canonical here."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pdfplumber
import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class DocumentLimitError(ValueError):
    """The input is valid enough to inspect but exceeds configured safety limits."""


class MalformedDocumentError(ValueError):
    """The input cannot be safely parsed as its declared format."""


class ScannedDocumentError(ValueError):
    """A PDF has no usable text layer and requires OCR/human review."""


@dataclass(frozen=True)
class IngestionLimits:
    max_file_bytes: int = 20 * 1024 * 1024
    max_uncompressed_bytes: int = 100 * 1024 * 1024
    max_zip_entries: int = 5_000
    max_sheets: int = 50
    max_cells: int = 200_000
    max_pages: int = 200
    max_page_chars: int = 200_000


@dataclass(frozen=True)
class CellEvidence:
    coordinate: str
    value: str | None
    original_value: str | None = None
    displayed_value: str | None = None
    formula: str | None = None
    merged_range: str | None = None


@dataclass(frozen=True)
class SheetEvidence:
    name: str
    cells: tuple[CellEvidence, ...]
    merged_ranges: tuple[str, ...]


@dataclass(frozen=True)
class WordEvidence:
    text: str
    x0: float
    top: float
    x1: float
    bottom: float


@dataclass(frozen=True)
class PageEvidence:
    page: int
    text: str
    words: tuple[WordEvidence, ...]
    tables: tuple[tuple[tuple[str | None, ...], ...], ...] = field(default_factory=tuple)


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    media_type: str
    sheets: tuple[SheetEvidence, ...] = field(default_factory=tuple)
    pages: tuple[PageEvidence, ...] = field(default_factory=tuple)


DEFAULT_LIMITS = IngestionLimits()


def _bounded_bytes(source: bytes | Path, limits: IngestionLimits) -> tuple[bytes, str]:
    if isinstance(source, Path):
        size = source.stat().st_size
        if size > limits.max_file_bytes:
            raise DocumentLimitError(f"File exceeds {limits.max_file_bytes} byte limit")
        return source.read_bytes(), source.name
    if len(source) > limits.max_file_bytes:
        raise DocumentLimitError(f"File exceeds {limits.max_file_bytes} byte limit")
    return source, "upload"


def parse_xlsx(source: bytes | Path, limits: IngestionLimits = DEFAULT_LIMITS) -> ParsedDocument:
    data, filename = _bounded_bytes(source, limits)
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_zip_entries:
                raise DocumentLimitError("Workbook contains too many ZIP entries")
            total = sum(entry.file_size for entry in entries)
            if total > limits.max_uncompressed_bytes:
                raise DocumentLimitError("Workbook expands beyond the uncompressed size limit")
            if any(entry.file_size > limits.max_uncompressed_bytes for entry in entries):
                raise DocumentLimitError("Workbook contains an oversized ZIP member")
    except (BadZipFile, OSError) as exc:
        raise MalformedDocumentError("Invalid XLSX container") from exc

    try:
        formulas = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
        values = load_workbook(BytesIO(data), read_only=False, data_only=True, keep_links=False)
    except Exception as exc:  # openpyxl exposes several format/XML exceptions
        raise MalformedDocumentError("XLSX content could not be parsed") from exc
    if len(formulas.worksheets) > limits.max_sheets:
        raise DocumentLimitError("Workbook contains too many sheets")

    sheets: list[SheetEvidence] = []
    seen_cells = 0
    try:
        for ws in formulas.worksheets:
            value_ws = values[ws.title]
            merged_ranges = tuple(str(item) for item in ws.merged_cells.ranges)
            merged_by_coordinate = {
                coordinate: merged
                for merged in merged_ranges
                for row in ws[merged]
                for cell in row
                for coordinate in [cell.coordinate]
            }
            cells: list[CellEvidence] = []
            for row in ws.iter_rows():
                for cell in row:
                    formula = cell.value if cell.data_type == "f" else None
                    raw = value_ws[cell.coordinate].value if formula else cell.value
                    if raw is None and formula is None and cell.coordinate not in merged_by_coordinate:
                        continue
                    seen_cells += 1
                    if seen_cells > limits.max_cells:
                        raise DocumentLimitError("Workbook contains too many populated cells")
                    cells.append(
                        CellEvidence(
                            coordinate=cell.coordinate,
                            value=None if raw is None else str(raw),
                            original_value=None if cell.value is None else str(cell.value),
                            displayed_value=None if raw is None else str(raw),
                            formula=None if formula is None else str(formula),
                            merged_range=merged_by_coordinate.get(cell.coordinate),
                        )
                    )
            sheets.append(SheetEvidence(ws.title, tuple(cells), merged_ranges))
    finally:
        formulas.close()
        values.close()
    return ParsedDocument(filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", tuple(sheets))


def parse_xls(source: bytes | Path, limits: IngestionLimits = DEFAULT_LIMITS) -> ParsedDocument:
    """Read legacy BIFF/OLE Excel workbooks with the same bounded evidence contract as XLSX."""
    data, filename = _bounded_bytes(source, limits)
    if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise MalformedDocumentError("Invalid XLS container")
    try:
        workbook = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    except (xlrd.XLRDError, OSError, ValueError) as exc:
        raise MalformedDocumentError("XLS content could not be parsed") from exc
    if workbook.nsheets > limits.max_sheets:
        raise DocumentLimitError("Workbook contains too many sheets")

    sheets: list[SheetEvidence] = []
    seen_cells = 0
    try:
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            if sheet.nrows * sheet.ncols > limits.max_cells:
                raise DocumentLimitError("Workbook contains too many cells")
            merged_ranges = tuple(
                f"{get_column_letter(left + 1)}{top + 1}:{get_column_letter(right)}{bottom}"
                for top, bottom, left, right in sheet.merged_cells
            )
            merged_by_coordinate = {
                f"{get_column_letter(column + 1)}{row + 1}": merged
                for merged, (top, bottom, left, right) in zip(merged_ranges, sheet.merged_cells)
                for row in range(top, bottom)
                for column in range(left, right)
            }
            cells: list[CellEvidence] = []
            for row in range(sheet.nrows):
                for column in range(sheet.ncols):
                    cell = sheet.cell(row, column)
                    coordinate = f"{get_column_letter(column + 1)}{row + 1}"
                    if cell.ctype == xlrd.XL_CELL_EMPTY and coordinate not in merged_by_coordinate:
                        continue
                    seen_cells += 1
                    if seen_cells > limits.max_cells:
                        raise DocumentLimitError("Workbook contains too many populated cells")
                    raw = None if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else str(cell.value)
                    cells.append(
                        CellEvidence(
                            coordinate=coordinate,
                            value=raw,
                            original_value=raw,
                            displayed_value=raw,
                            merged_range=merged_by_coordinate.get(coordinate),
                        )
                    )
            sheets.append(SheetEvidence(sheet.name, tuple(cells), merged_ranges))
    finally:
        workbook.release_resources()
    return ParsedDocument("upload" if filename == "upload" else filename, "application/vnd.ms-excel", tuple(sheets))


def parse_pdf(source: bytes | Path, limits: IngestionLimits = DEFAULT_LIMITS) -> ParsedDocument:
    data, filename = _bounded_bytes(source, limits)
    try:
        pdf = pdfplumber.open(BytesIO(data))
    except Exception as exc:
        raise MalformedDocumentError("PDF content could not be parsed") from exc
    pages: list[PageEvidence] = []
    total_text = 0
    try:
        if len(pdf.pages) > limits.max_pages:
            raise DocumentLimitError("PDF contains too many pages")
        for number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if len(text) > limits.max_page_chars:
                raise DocumentLimitError(f"PDF page {number} contains too much text")
            total_text += len(text.strip())
            words = tuple(
                WordEvidence(str(word["text"]), float(word["x0"]), float(word["top"]), float(word["x1"]), float(word["bottom"]))
                for word in page.extract_words()
            )
            tables = tuple(
                tuple(tuple(None if cell is None else str(cell) for cell in row) for row in table)
                for table in page.extract_tables()
            )
            pages.append(PageEvidence(number, text, words, tables))
    finally:
        pdf.close()
    if total_text == 0:
        raise ScannedDocumentError("No searchable text found. Run OCR or enter the document manually, then review the result.")
    return ParsedDocument(filename, "application/pdf", pages=tuple(pages))


def cell_range(sheet: str, row: int, start_column: int, end_column: int) -> str:
    """Build an auditable A1 range for deterministic table providers."""
    return f"'{sheet}'!{get_column_letter(start_column)}{row}:{get_column_letter(end_column)}{row}"
