from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from appraise_worker.extraction import FakeTableProvider, extract_draft
from appraise_worker.ingestion import (
    DocumentLimitError,
    IngestionLimits,
    MalformedDocumentError,
    ScannedDocumentError,
    parse_pdf,
    parse_xlsx,
)


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Specification"
    sheet.append(["Position", "Description", "Quantity", "Unit", "Price", "Total", "Currency"])
    sheet.append(["1", "Steel plate", 2, "kg", 3, "=C2*E2", "USD"])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Merged note"
    workbook.create_sheet("Terms")["A1"] = "Valid for 30 days"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes(*, with_text: bool = True, pages: int = 2) -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output)
    for index in range(pages):
        if with_text:
            pdf.drawString(72, 720, f"Line {index + 1}: steel plate 2 kg")
        pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_xlsx_preserves_sheets_formula_coordinates_and_merged_ranges() -> None:
    parsed = parse_xlsx(workbook_bytes())
    assert [sheet.name for sheet in parsed.sheets] == ["Specification", "Terms"]
    formula = next(cell for cell in parsed.sheets[0].cells if cell.coordinate == "F2")
    assert formula.formula == "=C2*E2"
    assert formula.value is None  # openpyxl does not calculate formula caches
    assert "A4:B4" in parsed.sheets[0].merged_ranges
    assert next(cell for cell in parsed.sheets[0].cells if cell.coordinate == "A4").merged_range == "A4:B4"


def test_xlsx_rejects_malformed_and_zip_expansion_limits() -> None:
    with pytest.raises(MalformedDocumentError):
        parse_xlsx(b"not-an-xlsx")
    bomb = BytesIO()
    with ZipFile(bomb, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"x" * 1_000)
    with pytest.raises(DocumentLimitError, match="uncompressed"):
        parse_xlsx(bomb.getvalue(), IngestionLimits(max_uncompressed_bytes=100))


def test_pdf_preserves_multi_page_text_and_word_coordinates() -> None:
    parsed = parse_pdf(pdf_bytes())
    assert len(parsed.pages) == 2
    assert parsed.pages[1].page == 2
    assert "steel plate" in parsed.pages[0].text
    assert parsed.pages[0].words[0].x0 >= 0


def test_digital_pdf_produces_reviewable_draft_with_page_provenance() -> None:
    draft = extract_draft(FakeTableProvider(), "specification", parse_pdf(pdf_bytes()))
    assert len(draft.lines) == 2
    assert draft.lines[1].source.page == 2
    assert draft.lines[0].description == "steel plate"


def test_pdf_rejects_scanned_and_page_limit() -> None:
    with pytest.raises(ScannedDocumentError, match="OCR"):
        parse_pdf(pdf_bytes(with_text=False))
    with pytest.raises(DocumentLimitError, match="too many pages"):
        parse_pdf(pdf_bytes(pages=2), IngestionLimits(max_pages=1))
