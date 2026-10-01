from io import BytesIO

from openpyxl import load_workbook

from appraise_worker.proposal_export import build_workbook


def fixture() -> dict:
    return {
        "seller_name": "Appraise Industrial",
        "seller_legal_name": "Appraise Industrial LLP",
        "client_name": "Barys Client",
        "project_name": "Barys repair",
        "proposal_number": "APP-2026-001",
        "version": 2,
        "currency": "KZT",
        "created_at": "2026-09-20T10:00:00Z",
        "validity_days": 30,
        "notes": "Prices include delivery.",
        "commercial_terms": "Payment within 10 days.",
        "subtotal": "500000.00000000",
        "markup_amount": "50000.00000000",
        "tax_amount": "66000.00000000",
        "grand_total": "616000.00000000",
        "lines": [{
            "position": "1.1", "description": "Pipe DN50", "quantity": "2", "unit": "m",
            "unit_price": "250000", "line_total": "500000",
        }],
    }


def test_proposal_workbook_is_styled_paginated_and_client_safe() -> None:
    content = build_workbook(fixture())
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Proposal"]
    assert content.startswith(b"PK")
    assert sheet.freeze_panes == "A9"
    assert sheet["F9"].value == 500000
    assert sheet["F14"].value == 616000
    assert sheet.print_area == "'Proposal'!$A$1:$F$19"
    assert "supplier_id" not in " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row)
